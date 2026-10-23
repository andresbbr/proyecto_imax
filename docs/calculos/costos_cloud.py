"""Actualiza los volúmenes del estimado de costos cloud (AWS vs Azure).

Copia docs/fuentes/estimado_cloud_imax.xlsx a docs/calculos/costos_cloud.xlsx y reemplaza
en la hoja Supuestos solo los volúmenes calculados en dimensionamiento.md (celdas B9, B10,
B13, B23 y B24). Los precios unitarios y las fórmulas no se tocan.

Después hay que abrir costos_cloud.xlsx en LibreOffice Calc o Excel, recalcular
(Ctrl+Mayús+F9 en LibreOffice) y guardar. Con el libro ya recalculado, el script
muestra la hoja Resumen:

    py docs/calculos/costos_cloud.py             # escribe los volúmenes
    py docs/calculos/costos_cloud.py --resumen   # muestra el Resumen recalculado

Requiere openpyxl.
"""

import pathlib
import re
import shutil
import sys

import openpyxl

CARPETA = pathlib.Path(__file__).parent
ORIGINAL = CARPETA.parent / "fuentes" / "estimado_cloud_imax.xlsx"
LIBRO = CARPETA / "costos_cloud.xlsx"
CELDAS = ["B9", "B10", "B13", "B23", "B24"]


def leer_volumenes():
    """Lee la tabla de volúmenes de dimensionamiento.md."""
    texto = (CARPETA / "dimensionamiento.md").read_text(encoding="utf-8")
    volumenes = {}
    for celda in CELDAS:
        m = re.search(rf"^\| {celda} \| [^|]+ \| ([\d,.]+) \|", texto, re.MULTILINE)
        valor = float(m.group(1).replace(",", ""))
        volumenes[celda] = int(valor) if valor.is_integer() else valor
    return volumenes


def escribir_volumenes():
    shutil.copy(ORIGINAL, LIBRO)
    wb = openpyxl.load_workbook(LIBRO)
    hoja = wb["Supuestos"]
    for celda, valor in leer_volumenes().items():
        print(f"Supuestos!{celda}: {hoja[celda].value} -> {valor}  ({hoja['A' + celda[1:]].value})")
        hoja[celda] = valor
    wb.save(LIBRO)
    print(f"\nEscrito {LIBRO}. Ábrelo en LibreOffice o Excel, recalcula y guárdalo.")


def mostrar_resumen():
    hoja = openpyxl.load_workbook(LIBRO, data_only=True)["Resumen"]
    if hoja["B11"].value is None:
        sys.exit("El libro no está recalculado: ábrelo en LibreOffice o Excel y guárdalo.")
    for fila in hoja.iter_rows(min_row=4, max_row=25, max_col=5):
        if fila[0].value is None:
            continue
        valores = [f"{c.value:,.6f}" if isinstance(c.value, float) and abs(c.value) < 0.01
                   else f"{c.value:,.2f}" if isinstance(c.value, (int, float)) else str(c.value or "")
                   for c in fila[1:]]
        print(f"{fila[0].value:<55} " + " ".join(f"{v:>20}" for v in valores))


if __name__ == "__main__":
    if "--resumen" in sys.argv:
        mostrar_resumen()
    else:
        escribir_volumenes()
