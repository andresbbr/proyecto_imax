"""Arma el documento técnico final en Word a partir de las secciones en Markdown.

Une 01_*.md a 11_*.md en orden y escribe documento_tecnico_imax.docx en esta carpeta.
Formato: Arial 11 en todo el texto (también tablas y listas), títulos en negrita sin
estilos de encabezado de Word, diagramas como imágenes.

Uso:
    py docs/documento/generar_docx.py

Requiere python-docx.
"""

import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CARPETA = pathlib.Path(__file__).parent
SALIDA = CARPETA / "documento_tecnico_imax.docx"
FUENTE = "Arial"
TAMANO = Pt(11)
RAYA_LARGA = "—"
MARGEN_CM = 2.0
ANCHO_UTIL_CM = 17.5  # carta (21.59 cm) menos los márgenes


def fijar_fuente(run, negrita=None, cursiva=None):
    run.font.name = FUENTE
    run.font.size = TAMANO
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FUENTE)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva


def texto_en_linea(parrafo, texto, negrita=False):
    """Escribe texto con **negrita**, *cursiva* y `código` en línea."""
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            fijar_fuente(parrafo.add_run(parte[2:-2]), negrita=True)
        elif parte.startswith("`") and parte.endswith("`"):
            fijar_fuente(parrafo.add_run(parte[1:-1]), negrita=negrita)
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            fijar_fuente(parrafo.add_run(parte[1:-1]), negrita=negrita, cursiva=True)
        else:
            fijar_fuente(parrafo.add_run(parte), negrita=negrita)


def titulo(doc, texto, nivel):
    """Título en negrita con el estilo Normal (no se usan los estilos de encabezado)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if nivel > 1 else 0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    fijar_fuente(p.add_run(texto.replace("**", "").replace("`", "")), negrita=True)


def anchos_de_columna(celdas, columnas):
    """Reparte el ancho útil según el largo del texto de cada columna, con un mínimo para que
    la palabra más larga no se corte."""
    def palabra_mas_larga(texto):
        return max((len(p) for p in texto.replace("**", "").replace("`", "").split()), default=0)

    cm_por_caracter, relleno, total = 0.2, 0.4, ANCHO_UTIL_CM
    minimos, deseados = [], []
    for j in range(columnas):
        textos = [f[j] for f in celdas if j < len(f)]
        cuerpo = min(max(palabra_mas_larga(t) for t in textos[1:] or [""]), 14) * cm_por_caracter
        encabezado = palabra_mas_larga(textos[0]) * cm_por_caracter * 1.2  # en negrita ocupa más
        minimos.append(max(cuerpo, encabezado) + relleno)
        deseados.append(max(min(max(len(t) for t in textos), 150) * cm_por_caracter + relleno, minimos[-1]))
    if sum(minimos) >= total:
        return [Cm(total * m / sum(minimos)) for m in minimos]
    extra = [d - m for d, m in zip(deseados, minimos)]
    sobrante = total - sum(minimos)
    return [Cm(m + (sobrante * e / sum(extra) if sum(extra) else sobrante / columnas))
            for m, e in zip(minimos, extra)]


def tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    celdas = [f for f in celdas if not all(re.fullmatch(r":?-{3,}:?", c) for c in f)]
    columnas = max(len(f) for f in celdas)
    t = doc.add_table(rows=len(celdas), cols=columnas)
    t.style = "Table Grid"
    t.autofit = False
    anchos = anchos_de_columna(celdas, columnas)
    for j, ancho in enumerate(anchos):
        t.columns[j].width = ancho
    for i, fila in enumerate(celdas):
        for j in range(columnas):
            celda = t.cell(i, j)
            celda.width = anchos[j]
            texto_en_linea(celda.paragraphs[0], fila[j] if j < len(fila) else "", negrita=(i == 0))
    # La fila de encabezado se repite si la tabla pasa a otra página
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph()


def imagen(doc, ruta_md, alternativo):
    ruta = (CARPETA / ruta_md).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(ANCHO_UTIL_CM))
    leyenda = doc.add_paragraph()
    leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fijar_fuente(leyenda.add_run(f"Figura: {alternativo}"), cursiva=True)


def item_de_lista(doc, texto, sangria, numero=None):
    p = doc.add_paragraph()
    formato = p.paragraph_format
    formato.left_indent = Cm(0.75 + 0.75 * sangria)
    formato.first_line_indent = Cm(-0.5)
    formato.space_after = Pt(3)
    fijar_fuente(p.add_run(f"{numero}. " if numero else ("• " if sangria == 0 else "◦ ")))
    texto_en_linea(p, texto)


def codigo(doc, lineas):
    for linea in lineas:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(0)
        fijar_fuente(p.add_run(linea))
    doc.add_paragraph()


def agregar_seccion(doc, texto):
    lineas = texto.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        sin_espacios = linea.strip()
        if sin_espacios.startswith("```"):
            fin = i + 1
            while not lineas[fin].strip().startswith("```"):
                fin += 1
            codigo(doc, lineas[i + 1:fin])
            i = fin + 1
            continue
        if sin_espacios.startswith("|"):
            fin = i
            while fin < len(lineas) and lineas[fin].strip().startswith("|"):
                fin += 1
            tabla(doc, lineas[i:fin])
            i = fin
            continue
        encabezado = re.match(r"(#{1,4}) (.*)", linea)
        imagen_md = re.match(r"!\[(.*)\]\((.*)\)", sin_espacios)
        vineta = re.match(r"(\s*)[-*] (.*)", linea)
        numerado = re.match(r"(\s*)(\d+)\. (.*)", linea)
        if encabezado:
            titulo(doc, encabezado.group(2), len(encabezado.group(1)))
        elif imagen_md:
            imagen(doc, imagen_md.group(2), imagen_md.group(1))
        elif vineta:
            item_de_lista(doc, vineta.group(2), len(vineta.group(1)) // 2)
        elif numerado:
            item_de_lista(doc, numerado.group(3), len(numerado.group(1)) // 2, numero=numerado.group(2))
        elif sin_espacios:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            texto_en_linea(p, sin_espacios)
        i += 1


def numero_de_pagina(seccion):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fijar_fuente(run)
    for tipo, texto in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tipo:
            campo = OxmlElement("w:fldChar")
            campo.set(qn("w:fldCharType"), tipo)
        else:
            campo = OxmlElement("w:instrText")
            campo.set(qn("xml:space"), "preserve")
            campo.text = texto
        run._element.append(campo)


def verificar(ruta):
    doc = Document(ruta)
    errores = []
    parrafos = list(doc.paragraphs) + [p for t in doc.tables for f in t.rows for c in f.cells for p in c.paragraphs]
    for p in parrafos:
        if p.style.name.startswith(("Heading", "Title")):
            errores.append(f"estilo de encabezado: {p.text[:40]}")
        if RAYA_LARGA in p.text:
            errores.append(f"raya larga: {p.text[:40]}")
        for run in p.runs:
            if run.text.strip() and (run.font.name != FUENTE or run.font.size != TAMANO):
                errores.append(f"fuente distinta de Arial 11: {run.text[:40]}")
    imagenes = len(doc.inline_shapes)
    return errores, imagenes, len(doc.tables)


def main():
    secciones = sorted(CARPETA.glob("[01][0-9]_*.md"))
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.font.size = TAMANO
    estilo.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FUENTE)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(MARGEN_CM)
    numero_de_pagina(doc.sections[0])
    doc.core_properties.title = "Licitación para el Ecosistema Digital de Cines IMAX"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""

    for n, ruta in enumerate(secciones):
        if n > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        agregar_seccion(doc, ruta.read_text(encoding="utf-8"))
    doc.save(SALIDA)

    errores, imagenes, tablas = verificar(SALIDA)
    print(f"Escrito {SALIDA.name}: {len(secciones)} secciones, {tablas} tablas, {imagenes} imágenes")
    for e in errores[:20]:
        print("  ERROR", e)
    sys.exit(1 if errores else 0)


if __name__ == "__main__":
    main()
